import re
from pathlib import Path

import mysql.connector
from openpyxl import load_workbook

ACTIVE_DISTRICTS = {
    'Pulau Ternate',
    'Ternate Selatan',
    'Ternate Tengah',
    'Ternate Utara',
}

BASE_DIR = Path(__file__).parent


def number(value):
    return float(value or 0)


connection = mysql.connector.connect(
    host='127.0.0.1', port=3306, user='root', password='', database='gis_ternate'
)
cursor = connection.cursor()
cursor.execute('''
CREATE TABLE IF NOT EXISTS crop_production (
  id INT AUTO_INCREMENT PRIMARY KEY,
    district_name VARCHAR(40) NOT NULL,
    commodity VARCHAR(100) NOT NULL,
  year SMALLINT NOT NULL,
  production_ton DECIMAL(12,3) NOT NULL,
  source_file VARCHAR(255) NOT NULL,
  UNIQUE KEY crop_record (district_name, commodity, year)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
''')
cursor.execute('''
CREATE TABLE IF NOT EXISTS horticulture_production (
  id INT AUTO_INCREMENT PRIMARY KEY,
  commodity VARCHAR(150) NOT NULL,
  year SMALLINT NOT NULL,
  production_ton DECIMAL(12,3) NOT NULL,
  source_file VARCHAR(255) NOT NULL,
  UNIQUE KEY horticulture_record (commodity, year)
)
''')

food_path = BASE_DIR / 'data produksi tanaman pangan.xlsx'
hort_path = BASE_DIR / 'data produksi Hortikultura.xlsx'
if not food_path.exists() or not hort_path.exists():
    missing_files = [path.name for path in (food_path, hort_path) if not path.exists()]
    print('Import dibatalkan; file sumber tidak ditemukan:', ', '.join(missing_files))
    raise SystemExit(0)

food_workbook = load_workbook(food_path, read_only=True, data_only=True)
latest_totals = {district: 0 for district in ACTIVE_DISTRICTS}
food_rows = []
food_years = [int(sheet.title) for sheet in food_workbook.worksheets if str(sheet.title).isdigit()]
latest_food_year = max(food_years)

for sheet in food_workbook.worksheets:
    year = int(sheet.title)
    commodity = None
    for row in sheet.iter_rows(values_only=True):
        title = str(row[0] or '')
        match = re.search(r'Produksi(?:.*?)\s+(Jagung|Ubi Kayu|Ubi Jalar|Kacang Tanah|kedelai)', title, re.I)
        if match:
            commodity = match.group(1).title()
            continue
        district = row[1] if len(row) > 1 else None
        production = row[4] if len(row) > 4 else None
        if district in ACTIVE_DISTRICTS and isinstance(production, (int, float)) and commodity:
            food_rows.append((district, commodity, year, number(production), food_path.name))
            if year == latest_food_year:
                latest_totals[district] += number(production)

for row in food_rows:
    cursor.execute('''
      INSERT INTO crop_production (district_name, commodity, year, production_ton, source_file)
      VALUES (%s, %s, %s, %s, %s)
      ON DUPLICATE KEY UPDATE production_ton = VALUES(production_ton), source_file = VALUES(source_file)
    ''', row)

hort_workbook = load_workbook(hort_path, read_only=True, data_only=True)
hort_rows = []
sheet = hort_workbook.active
for row in sheet.iter_rows(min_row=9, values_only=True):
    commodity = row[2] if len(row) > 2 else None
    if not commodity:
        continue
    for index, year in enumerate(range(2021, 2026), start=3):
        production = row[index] if len(row) > index else None
        if isinstance(production, (int, float)):
            hort_rows.append((str(commodity).strip(), year, number(production), hort_path.name))

for row in hort_rows:
    cursor.execute('''
      INSERT INTO horticulture_production (commodity, year, production_ton, source_file)
      VALUES (%s, %s, %s, %s)
      ON DUPLICATE KEY UPDATE production_ton = VALUES(production_ton), source_file = VALUES(source_file)
    ''', row)

for district, total in latest_totals.items():
    cursor.execute(
        "UPDATE districts SET production_pertanian = %s, production_total = production_ikan + %s, "
        "commodity = %s, type_name = 'Hasil Pertanian' WHERE name = %s",
        (total, total, f'Tanaman pangan {latest_food_year}', district),
    )

connection.commit()
print('Data pangan diimpor:', len(food_rows), 'baris')
print('Data hortikultura diimpor:', len(hort_rows), 'baris')
print(f'Produksi pertanian {latest_food_year} per zona:', latest_totals)
cursor.close()
connection.close()
