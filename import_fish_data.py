import re
from pathlib import Path

from openpyxl import load_workbook


FISHERY_LOCATIONS = {
    'DUFA DUFA': 'Ternate Utara',
}


def has_location(text, keyword):
    normalized_text = re.sub(r'[^A-Z0-9]', '', text.upper())
    normalized_keyword = re.sub(r'[^A-Z0-9]', '', keyword.upper())
    return normalized_keyword in normalized_text


def parse_kg(value):
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or '').strip().lower()
    if not text or text in {'-', '0'}:
        return 0.0
    digits = re.sub(r'[^0-9]', '', text.replace('kg', ''))
    return float(digits or 0)


def find_fishery_files(folder):
    pdf_stems = {path.stem.casefold() for path in folder.glob('*.pdf')}
    return sorted(
        path for path in folder.glob('*.xlsx')
        if path.stem.casefold() in pdf_stems
        and 'PRODUKSI' in path.stem.upper()
    )


def read_workbook(path):
    match = re.search(r'(20\d{2})', path.name)
    year = int(match.group(1)) if match else None
    workbook = load_workbook(path, read_only=True, data_only=True)
    location = next(
        (district for keyword, district in FISHERY_LOCATIONS.items() if has_location(path.stem, keyword)),
        None,
    )
    if location is None:
        for sheet in workbook.worksheets:
            title = str(next(sheet.iter_rows(values_only=True))[0] or '').upper()
            location = next(
                (district for keyword, district in FISHERY_LOCATIONS.items() if has_location(title, keyword)),
                None,
            )
            if location:
                break
    if year is None or location is None:
        return []

    rows = []
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        header_index = next(
            (
                index for index, row in enumerate(values[:-1])
                if any(str(value or '').strip().upper() == 'BULAN' for value in row)
                and any('JENIS' in str(value or '').upper() for value in row)
            ),
            None,
        )
        if header_index is None:
            continue
        species = [
            str(value).replace('\n', ' ').strip()
            for value in values[header_index + 1][2:]
            if value and 'TOTAL' not in str(value).upper()
        ]
        for row in values[header_index + 2:]:
            month = str(row[1] or '').strip() if len(row) > 1 else ''
            if not month or month.upper() == 'TOTAL':
                continue
            for index, commodity in enumerate(species, start=2):
                if index < len(row):
                    production_ton = parse_kg(row[index]) / 1000
                    if production_ton:
                        rows.append((location, commodity, year, production_ton, path.name))
    return rows


def main():
    folder = Path(__file__).parent
    parsed_rows = [row for path in find_fishery_files(folder) for row in read_workbook(path)]
    totals = {}
    for location, commodity, year, production_ton, source_file in parsed_rows:
        key = (location, commodity, year)
        previous_total = totals.get(key, (0, source_file))[0]
        totals[key] = (previous_total + production_ton, source_file)
    fish_rows = [(*key, production_ton, source_file) for key, (production_ton, source_file) in totals.items()]
    if not fish_rows:
        raise SystemExit('Tidak ditemukan XLSX perikanan yang memiliki pasangan PDF.')

    from app import get_db_connection

    connection = get_db_connection()
    cursor = connection.cursor()
    cursor.execute('''
      CREATE TABLE IF NOT EXISTS fish_production (
        id INT AUTO_INCREMENT PRIMARY KEY,
                district_name VARCHAR(40) NOT NULL,
                commodity VARCHAR(100) NOT NULL,
        year SMALLINT NOT NULL,
        production_ton DECIMAL(12,3) NOT NULL,
        source_file VARCHAR(255) NOT NULL,
        UNIQUE KEY fish_record (district_name, commodity, year)
    ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    ''')
    for row in fish_rows:
        cursor.execute('''
          INSERT INTO fish_production
            (district_name, commodity, year, production_ton, source_file)
          VALUES (%s, %s, %s, %s, %s)
          ON DUPLICATE KEY UPDATE
                        production_ton = VALUES(production_ton),
            source_file = VALUES(source_file)
        ''', row)

    latest_year = max(row[2] for row in fish_rows)
    cursor.execute('''
      UPDATE districts district
      LEFT JOIN (
        SELECT district_name, SUM(production_ton) AS production
        FROM fish_production
        WHERE year = %s
        GROUP BY district_name
      ) fish ON fish.district_name = district.name
      SET district.production_ikan = COALESCE(fish.production, 0),
          district.production_total = COALESCE(fish.production, 0) + district.production_pertanian
    ''', (latest_year,))
    connection.commit()
    print('Data perikanan diimpor:', len(fish_rows), 'baris')
    print('Tahun terbaru:', latest_year)
    cursor.close()
    connection.close()


if __name__ == '__main__':
    main()